"""Exportación e importación de planillas de comprobantes."""

from __future__ import annotations

import json
import re
from base64 import urlsafe_b64decode, urlsafe_b64encode
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from typing import Any
from uuid import UUID, uuid4

from vrf.adapters.outbound.ocr.afip import parse_importe
from vrf.application.use_cases.comprobantes import (
    CrearComprobante,
    DatosAlta,
    ListarComprobantes,
    formatear_numero_comprobante,
)
from vrf.domain.entities import Proveedor, Rubro, Usuario
from vrf.domain.enums import AmbitoRubro, Clasificacion, ClaseRubro, TipoAfip
from vrf.domain.exceptions import Conflict, InvalidInput
from vrf.domain.ports.planilla import EscritorPlanilla
from vrf.domain.ports.repositories import (
    ClienteRepository,
    ComprobanteRepository,
    ObraRepository,
    ProveedorRepository,
    RubroRepository,
    TipoPagoRepository,
    UnitOfWork,
)
from vrf.domain.vos import dinero


def _money(valor: Any) -> Decimal | None:
    """Acepta celdas Excel (int/float/Decimal) y textos AR ($ 1.234,56)."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, Decimal):
        return dinero(valor)
    if isinstance(valor, (int, float)):
        return dinero(valor)
    s = str(valor).strip()
    parsed = parse_importe(s)
    if parsed is not None:
        return dinero(parsed)
    try:
        return dinero(s.replace(" ", "").replace("$", "").replace(",", "."))
    except Exception:
        return None


def _fmt_fecha(valor: date | None) -> str:
    if valor is None:
        return ""
    return valor.strftime("%d/%m/%Y")


def _fmt_money(valor: Decimal | str | None) -> str:
    if valor is None:
        return ""
    n = Decimal(str(valor))
    # Formato argentino: 1.234,56
    s = f"{n:.2f}"
    entero, dec = s.split(".")
    partes: list[str] = []
    while entero:
        partes.append(entero[-3:])
        entero = entero[:-3]
    return f"{'.'.join(reversed(partes))},{dec}"


def _celda(c, nombres: dict[str, dict[str, str]], col: str) -> str:
    prov = nombres.get("proveedores", {})
    obras = nombres.get("obras", {})
    rubros = nombres.get("rubros", {})
    tipos = nombres.get("tipos", {})
    clientes = nombres.get("clientes", {})
    if col == "Fecha":
        return _fmt_fecha(c.fecha)
    if col == "Proveedor":
        return prov.get(str(c.proveedor_id or ""), "")
    if col == "Factura":
        return c.numero_comprobante or formatear_numero_comprobante(None, c.punto_venta, c.numero) or ""
    if col == "Clasificación":
        return c.clasificacion.value
    if col == "Rubro/Tipo":
        if c.clasificacion == Clasificacion.PAGO:
            return tipos.get(str(c.tipo_pago_id or ""), "")
        return rubros.get(str(c.rubro_id or ""), "")
    if col in ("Obra", "Detalle"):
        return obras.get(str(c.obra_id or ""), "")
    if col == "Cliente":
        return clientes.get(str(c.obra_id or ""), "")
    if col == "Gasto":
        return rubros.get(str(c.rubro_id or ""), "")
    if col == "Importe":
        return _fmt_money(c.neto_21 + c.neto_105 + c.neto_27 + c.no_gravado)
    if col == "IVA":
        return _fmt_money(c.iva_21 + c.iva_105 + c.iva_27)
    if col == "Total":
        return _fmt_money(c.total)
    if col == "Estado":
        return c.estado_pago.value
    if col == "Tipo AFIP":
        return c.tipo_afip.value
    if col == "Punto de venta":
        return "" if c.punto_venta is None else str(c.punto_venta)
    if col == "Número":
        return "" if c.numero is None else str(c.numero)
    if col == "CUIT emisor":
        return c.cuit_emisor or ""
    mapa = {
        "Neto 21%": c.neto_21,
        "IVA 21%": c.iva_21,
        "Neto 10,5%": c.neto_105,
        "IVA 10,5%": c.iva_105,
        "Neto 27%": c.neto_27,
        "IVA 27%": c.iva_27,
        "No gravado": c.no_gravado,
        "Percepción IVA": c.percep_iva,
        "Percepción IIBB": c.percep_iibb,
    }
    if col in mapa:
        return _fmt_money(mapa[col])
    return ""


class ExportarRegistros:
    def __init__(
        self,
        listar: ListarComprobantes,
        escritor: EscritorPlanilla,
        obras: ObraRepository,
        proveedores: ProveedorRepository,
        rubros: RubroRepository,
        tipos: TipoPagoRepository,
        clientes: ClienteRepository,
    ) -> None:
        self._listar = listar
        self._escritor = escritor
        self._obras = obras
        self._proveedores = proveedores
        self._rubros = rubros
        self._tipos = tipos
        self._clientes = clientes

    def execute(
        self,
        actor: Usuario,
        columnas: list[str],
        **filtros: Any,
    ) -> bytes:
        if not columnas:
            raise InvalidInput("Indicá columnas a exportar")
        comps = self._listar.execute(actor, limit=2000, **filtros)
        nombres = {
            "proveedores": {},
            "obras": {},
            "rubros": {},
            "tipos": {},
            "clientes": {},
        }
        clientes = {str(c.id): c.nombre for c in self._clientes.list_by_grupo(actor.grupo_id)}
        for p in self._proveedores.list_by_grupo(actor.grupo_id):
            nombres["proveedores"][str(p.id)] = p.razon_social
        for o in self._obras.list_by_grupo(actor.grupo_id):
            nombres["obras"][str(o.id)] = o.nombre
            nombres["clientes"][str(o.id)] = clientes.get(str(o.cliente_id), "")
        if filtros.get("empresa_id"):
            eid = filtros["empresa_id"]
            for r in self._rubros.list_by_empresa(eid):
                nombres["rubros"][str(r.id)] = r.nombre
            for t in self._tipos.list_by_empresa(eid):
                nombres["tipos"][str(t.id)] = t.nombre
        filas = [[_celda(c, nombres, col) for col in columnas] for c in comps]
        return self._escritor.xlsx(columnas, filas)


@dataclass
class FilaImport:
    fila: int
    empresa_id: str
    clasificacion: str
    fecha: str
    proveedor: str | None
    factura: str | None
    importe: str | None
    iva: str | None
    total: str | None
    gasto: str | None
    detalle: str | None
    estado: str | None
    cuit: str | None
    tipo_afip: str | None
    neto_21: str | None
    iva_21: str | None
    neto_105: str | None
    iva_105: str | None
    neto_27: str | None
    iva_27: str | None
    no_gravado: str | None
    percep_iva: str | None
    percep_iibb: str | None
    obra: str | None
    rubro: str | None


_ENCABEZADOS = {
    "fecha": "fecha",
    "proveedor": "proveedor",
    "factura": "factura",
    "numero factura": "factura",
    "número": "numero_afip",
    "numero": "numero_afip",
    "punto de venta": "punto_venta",
    "tipo afip": "tipo_afip_col",
    "importe": "importe",
    "gasto": "gasto",
    "iva": "iva",
    "total": "total",
    "estado": "estado",
    "detalle": "detalle",
    "c.u.i.t.": "cuit",
    "cuit": "cuit",
    "cuit emisor": "cuit",
    'factura. "c"': "neto_c",
    'factura. "a"': "neto_a",
    "iva 0,21": "iva_21",
    "0.105": "neto_105",
    "iva 0,27": "iva_27",
    "iva 10,5": "iva_105",
    "iva 21%": "iva_21",
    "neto 21%": "neto_21",
    "neto 10,5%": "neto_105",
    "iva 10,5%": "iva_105",
    "neto 27%": "neto_27",
    "iva 27%": "iva_27",
    "no.g": "no_gravado",
    "no gravado": "no_gravado",
    "p.iva": "percep_iva",
    "percepción iva": "percep_iva",
    "p.ib": "percep_iibb",
    "percepción iibb": "percep_iibb",
    "obra": "obra",
    "empresa": "empresa",
    "rubro": "rubro",
    "rubro/tipo": "gasto",
    "clasificación": "clasificacion_col",
    "clasificacion": "clasificacion_col",
}


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _parse_fecha(valor: Any) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(s[:19] if " " in s and "T" not in s else s, fmt).date()
        except ValueError:
            continue
    # Español: 01-jul-26
    meses = {
        "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
        "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
    }
    m = re.match(r"^(\d{1,2})[-/]([A-Za-z]{3})[-/](\d{2,4})$", s)
    if m:
        mes = meses.get(m.group(2).lower()[:3])
        if mes:
            anio = int(m.group(3))
            if anio < 100:
                anio += 2000
            try:
                return date(anio, mes, int(m.group(1)))
            except ValueError:
                return None
    return None


def _parse_numero_factura(texto: str | None) -> tuple[TipoAfip, int | None, int | None, str | None]:
    if not texto or not str(texto).strip():
        return TipoAfip.NINGUNO, None, None, None
    raw = str(texto).strip()
    if raw.upper() in ("SF", "S/F", "SIN FACTURA", "-"):
        return TipoAfip.NINGUNO, None, None, raw.upper() if raw.upper() == "SF" else "SF"
    m = re.match(r"^(\d{1,5})\s*[-/]\s*(\d{1,8})$", raw)
    if m:
        return TipoAfip.C, int(m.group(1)), int(m.group(2)), formatear_numero_comprobante(None, int(m.group(1)), int(m.group(2)))
    if re.fullmatch(r"\d+", raw):
        return TipoAfip.NINGUNO, None, int(raw), raw
    return TipoAfip.NINGUNO, None, None, raw


def leer_filas_xlsx(data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook
    from io import BytesIO

    wb = load_workbook(BytesIO(data), data_only=True)
    # Preferir hoja MES (compras) si existe
    ws = wb["MES"] if "MES" in wb.sheetnames else wb.active
    filas_raw = list(ws.iter_rows(values_only=True))
    encabezado_idx = None
    headers_norm: list[str] = []
    for i, fila in enumerate(filas_raw[:15]):
        celdas = [_norm_header(c) for c in fila if c is not None and str(c).strip()]
        if not celdas:
            continue
        if any(h in _ENCABEZADOS for h in celdas) and ("fecha" in celdas or "proveedor" in celdas):
            encabezado_idx = i
            headers_norm = [_norm_header(c) for c in fila]
            break
    if encabezado_idx is None:
        raise InvalidInput("No se encontraron encabezados reconocibles en el Excel")
    out: list[dict[str, Any]] = []
    for offset, fila in enumerate(filas_raw[encabezado_idx + 1 :], start=encabezado_idx + 2):
        if not any(c is not None and str(c).strip() for c in fila):
            continue
        # Totales / resumen
        texto = " ".join(str(c) for c in fila if c is not None).lower()
        if "total " in texto and not str(fila[0] or "").strip():
            continue
        row: dict[str, Any] = {"_fila": offset}
        for idx, h in enumerate(headers_norm):
            key = _ENCABEZADOS.get(h)
            if not key or idx >= len(fila):
                continue
            row[key] = fila[idx]
        out.append(row)
    return headers_norm, out


class PreviaImportacion:
    def __init__(
        self,
        comps: ComprobanteRepository,
        proveedores: ProveedorRepository,
        rubros: RubroRepository,
        tipos: TipoPagoRepository,
        obras: ObraRepository,
        secret: str,
    ) -> None:
        self._comps = comps
        self._proveedores = proveedores
        self._rubros = rubros
        self._tipos = tipos
        self._obras = obras
        self._secret = secret

    def _firmar(self, payload: dict) -> str:
        raw = json.dumps(payload, separators=(",", ":"), ensure_ascii=False).encode()
        sig = sha256(self._secret.encode() + raw).hexdigest()[:32]
        return urlsafe_b64encode(raw).decode() + "." + sig

    def _verificar(self, token: str) -> dict:
        try:
            b64, sig = token.split(".", 1)
            raw = urlsafe_b64decode(b64.encode())
            esperado = sha256(self._secret.encode() + raw).hexdigest()[:32]
            if sig != esperado:
                raise InvalidInput("Token de importación inválido")
            return json.loads(raw.decode())
        except (ValueError, json.JSONDecodeError) as exc:
            raise InvalidInput("Token de importación inválido") from exc

    def execute(
        self,
        actor: Usuario,
        data: bytes,
        empresa_id: UUID,
        clasificacion_default: str | None = None,
    ) -> dict:
        _, filas = leer_filas_xlsx(data)
        errores: list[dict] = []
        ok: list[dict] = []
        proveedores = {p.razon_social.strip().upper(): p for p in self._proveedores.list_by_grupo(actor.grupo_id)}
        rubros = {r.nombre.strip().upper(): r for r in self._rubros.list_by_empresa(empresa_id)}
        tipos = {t.nombre.strip().upper(): t for t in self._tipos.list_by_empresa(empresa_id)}
        obras = {o.nombre.strip().upper(): o for o in self._obras.list_by_grupo(actor.grupo_id)}

        for row in filas:
            nfila = int(row["_fila"])
            fecha = _parse_fecha(row.get("fecha"))
            if fecha is None:
                errores.append({"fila": nfila, "columna": "Fecha", "mensaje": "Fecha inválida o vacía", "comprobanteId": None})
                continue
            factura_raw = row.get("factura")
            tipo_afip, pto, nro, nro_txt = _parse_numero_factura(
                None if factura_raw is None else str(factura_raw)
            )
            # Columnas separadas (export Libro IVA / registros)
            if row.get("punto_venta") not in (None, "") and row.get("numero_afip") not in (None, ""):
                try:
                    pto = int(re.sub(r"\D", "", str(row["punto_venta"])) or 0) or pto
                    nro = int(re.sub(r"\D", "", str(row["numero_afip"])) or 0) or nro
                    if nro_txt is None and pto is not None and nro is not None:
                        nro_txt = formatear_numero_comprobante(None, pto, nro)
                except ValueError:
                    pass
            if row.get("tipo_afip_col") not in (None, ""):
                try:
                    tipo_afip = TipoAfip(str(row["tipo_afip_col"]).strip().upper())
                except ValueError:
                    pass
            # Detectar A/C desde columnas de compras
            if row.get("neto_a") not in (None, ""):
                tipo_afip = TipoAfip.A
            elif row.get("neto_c") not in (None, ""):
                tipo_afip = TipoAfip.C

            clasif = (clasificacion_default or "costo").lower()
            if row.get("clasificacion_col") not in (None, ""):
                cand = str(row["clasificacion_col"]).strip().lower()
                if cand in ("compra", "pago", "costo"):
                    clasif = cand
            if row.get("gasto") and not row.get("iva") and clasificacion_default is None and row.get("clasificacion_col") in (None, ""):
                clasif = "costo"
            if row.get("detalle") is not None and "estado" in row and clasificacion_default is None and row.get("clasificacion_col") in (None, ""):
                clasif = "pago"
            if (
                clasificacion_default is None
                and row.get("clasificacion_col") in (None, "")
                and (row.get("neto_a") is not None or row.get("neto_c") is not None or row.get("cuit"))
            ):
                clasif = "compra"

            proveedor_id = None
            cuit = None
            if row.get("cuit"):
                digitos = re.sub(r"\D", "", str(row["cuit"]))
                if digitos:
                    cuit = digitos
            nombre_prov = str(row.get("proveedor") or "").strip()
            if nombre_prov:
                p = proveedores.get(nombre_prov.upper())
                if p is None and cuit:
                    # se crea en confirmación
                    pass
                elif p is None:
                    # crear al confirmar
                    pass
                else:
                    proveedor_id = str(p.id)
                    cuit = p.cuit or cuit

            total = _money(row.get("total"))
            importe = _money(row.get("importe"))
            iva_total = _money(row.get("iva"))
            neto_21 = _money(row.get("neto_a") if row.get("neto_a") not in (None, "") else row.get("neto_21"))
            if neto_21 is None:
                neto_21 = _money(row.get("neto_c"))
            iva_21 = _money(row.get("iva_21"))
            neto_105 = _money(row.get("neto_105")) or Decimal("0")
            iva_105 = _money(row.get("iva_105")) or Decimal("0")
            neto_27 = _money(row.get("neto_27")) or Decimal("0")
            iva_27 = _money(row.get("iva_27")) or Decimal("0")
            no_gravado = _money(row.get("no_gravado")) or Decimal("0")
            percep_iva = _money(row.get("percep_iva")) or Decimal("0")
            percep_iibb = _money(row.get("percep_iibb")) or Decimal("0")
            if clasif == "pago":
                if total is None:
                    total = (importe or Decimal("0")) + (iva_total or Decimal("0"))
                if importe is None:
                    importe = total - (iva_total or Decimal("0"))
            if clasif == "costo":
                if total is None:
                    total = importe
                if importe is None:
                    importe = total
            if clasif == "compra":
                if neto_21 is None and importe is not None:
                    neto_21 = importe
                if iva_21 is None and iva_total is not None:
                    iva_21 = iva_total
                partes = (
                    (neto_21 or Decimal("0"))
                    + (iva_21 or Decimal("0"))
                    + neto_105
                    + iva_105
                    + neto_27
                    + iva_27
                    + no_gravado
                    + percep_iva
                    + percep_iibb
                )
                if total is None:
                    total = partes if partes else None
            if total is None:
                errores.append({"fila": nfila, "columna": "Total", "mensaje": "Falta total/importe", "comprobanteId": None})
                continue

            # Duplicado AFIP si aplica
            if tipo_afip in (TipoAfip.A, TipoAfip.B, TipoAfip.C) and cuit and pto is not None and nro is not None:
                dup = self._comps.get_by_clave_afip(cuit, tipo_afip.value, pto, nro)
                if dup:
                    errores.append(
                        {
                            "fila": nfila,
                            "columna": "Factura",
                            "mensaje": f"Comprobante duplicado ({tipo_afip.value} {pto}-{nro})",
                            "comprobanteId": str(dup.id),
                        }
                    )
                    continue

            rubro_id = None
            rubro_nombre = None
            tipo_pago_id = None
            gasto = str(row.get("gasto") or row.get("rubro") or "").strip()
            if clasif in ("compra", "costo"):
                # COMPRAS.xlsx a veces no trae rubro o trae uno nuevo → fallback "Importación Libro IVA"
                if not gasto:
                    gasto = "Importación Libro IVA"
                r = rubros.get(gasto.upper())
                if r is None:
                    rubro_nombre = gasto
                else:
                    rubro_id = str(r.id)
            if clasif == "pago":
                gasto_tipo = gasto
                if not gasto_tipo and nombre_prov and "flete" in nombre_prov.lower():
                    gasto_tipo = "Flete/transporte"
                if gasto_tipo:
                    t = tipos.get(gasto_tipo.upper())
                    if t is None:
                        errores.append(
                            {
                                "fila": nfila,
                                "columna": "Gasto",
                                "mensaje": f"Tipo de pago desconocido: {gasto_tipo}",
                                "comprobanteId": None,
                            }
                        )
                        continue
                    tipo_pago_id = str(t.id)
                elif tipos:
                    tipo_pago_id = str(next(iter(tipos.values())).id)
                else:
                    errores.append(
                        {
                            "fila": nfila,
                            "columna": "Tipo",
                            "mensaje": "No hay tipos de pago en la empresa",
                            "comprobanteId": None,
                        }
                    )
                    continue

            obra_id = None
            detalle = str(row.get("detalle") or row.get("obra") or "").strip()
            if detalle:
                o = obras.get(detalle.upper())
                # Obra desconocida: se importa sin imputar (no bloquea el Libro IVA)
                if o is not None:
                    obra_id = str(o.id)

            ya_pagada = str(row.get("estado") or "").strip().lower() in (
                "comprobado",
                "comprobado_pagado",
                "pagado",
            )

            ok.append(
                {
                    "fila": nfila,
                    "empresa_id": str(empresa_id),
                    "obra_id": obra_id,
                    "proveedor_nombre": nombre_prov or None,
                    "proveedor_id": proveedor_id,
                    "cuit": cuit,
                    "clasificacion": clasif,
                    "rubro_id": rubro_id,
                    "rubro_nombre": rubro_nombre,
                    "tipo_pago_id": tipo_pago_id,
                    "tipo_afip": tipo_afip.value,
                    "punto_venta": pto,
                    "numero": nro,
                    "numero_comprobante": nro_txt,
                    "fecha": fecha.isoformat(),
                    "neto_21": str(neto_21 or (importe if clasif != "pago" else Decimal("0")) or Decimal("0")),
                    "iva_21": str(iva_21 or (iva_total if clasif == "pago" else Decimal("0")) or Decimal("0")),
                    "neto_105": str(neto_105),
                    "iva_105": str(iva_105),
                    "neto_27": str(neto_27),
                    "iva_27": str(iva_27),
                    "no_gravado": str(no_gravado),
                    "percep_iva": str(percep_iva),
                    "percep_iibb": str(percep_iibb),
                    "total": str(total),
                    "ya_pagada": ya_pagada,
                    "nota": f"import fila {nfila}",
                }
            )

        token = self._firmar({"empresa_id": str(empresa_id), "filas": ok}) if ok and not errores else ""
        # Política: no grabar si hay errores
        if errores:
            token = ""
        return {"token": token or "INVALID", "filas_ok": len(ok) if not errores else 0, "errores": errores, "_filas": ok if not errores else []}


class ConfirmarImportacion:
    def __init__(
        self,
        previa: PreviaImportacion,
        crear: CrearComprobante,
        proveedores: ProveedorRepository,
        rubros: RubroRepository,
        uow: UnitOfWork,
        secret: str,
    ) -> None:
        self._previa = previa
        self._crear = crear
        self._proveedores = proveedores
        self._rubros = rubros
        self._uow = uow
        self._secret = secret

    def execute(self, actor: Usuario, token: str) -> dict:
        payload = self._previa._verificar(token)
        creados = 0
        rubros_cache: dict[str, UUID] = {}
        try:
            for fila in payload.get("filas", []):
                proveedor_id = UUID(fila["proveedor_id"]) if fila.get("proveedor_id") else None
                if proveedor_id is None and (fila.get("proveedor_nombre") or fila.get("cuit")):
                    nombre = (fila.get("proveedor_nombre") or fila.get("cuit") or "Proveedor").strip()
                    cuit = fila.get("cuit")
                    existente = None
                    if cuit:
                        existente = self._proveedores.get_by_cuit(cuit)
                    if existente:
                        proveedor_id = existente.id
                    else:
                        item = Proveedor(
                            id=uuid4(),
                            grupo_id=actor.grupo_id,
                            razon_social=nombre,
                            cuit=cuit,
                        )
                        self._proveedores.add(item)
                        proveedor_id = item.id

                rubro_id = UUID(fila["rubro_id"]) if fila.get("rubro_id") else None
                if rubro_id is None and fila.get("rubro_nombre"):
                    nombre_r = str(fila["rubro_nombre"]).strip()
                    key = f"{fila['empresa_id']}:{nombre_r.upper()}"
                    if key in rubros_cache:
                        rubro_id = rubros_cache[key]
                    else:
                        eid = UUID(fila["empresa_id"])
                        hallado = next(
                            (
                                r
                                for r in self._rubros.list_by_empresa(eid)
                                if r.nombre.strip().upper() == nombre_r.upper()
                            ),
                            None,
                        )
                        if hallado:
                            rubro_id = hallado.id
                        else:
                            nuevo = Rubro(
                                id=uuid4(),
                                empresa_id=eid,
                                nombre=nombre_r,
                                clase=ClaseRubro.COMPRA_COSTO,
                                ambito=AmbitoRubro.AMBOS,
                            )
                            self._rubros.add(nuevo)
                            rubro_id = nuevo.id
                        rubros_cache[key] = rubro_id

                try:
                    self._crear.execute(
                        actor,
                        DatosAlta(
                            empresa_id=UUID(fila["empresa_id"]),
                            obra_id=UUID(fila["obra_id"]) if fila.get("obra_id") else None,
                            proveedor_id=proveedor_id,
                            clasificacion=Clasificacion(fila["clasificacion"]),
                            rubro_id=rubro_id,
                            tipo_pago_id=UUID(fila["tipo_pago_id"]) if fila.get("tipo_pago_id") else None,
                            tipo_afip=TipoAfip(fila["tipo_afip"]),
                            punto_venta=fila.get("punto_venta"),
                            numero=fila.get("numero"),
                            fecha=date.fromisoformat(fila["fecha"]),
                            neto_21=Decimal(fila["neto_21"]),
                            iva_21=Decimal(fila["iva_21"]),
                            neto_105=Decimal(fila["neto_105"]),
                            iva_105=Decimal(fila["iva_105"]),
                            neto_27=Decimal(fila["neto_27"]),
                            iva_27=Decimal(fila["iva_27"]),
                            no_gravado=Decimal(fila["no_gravado"]),
                            percep_iva=Decimal(fila["percep_iva"]),
                            percep_iibb=Decimal(fila["percep_iibb"]),
                            total=Decimal(fila["total"]),
                            ya_pagada=bool(fila.get("ya_pagada")),
                            nota=fila.get("nota"),
                            extra={},
                            numero_comprobante=fila.get("numero_comprobante"),
                        ),
                        commit=False,
                    )
                    creados += 1
                except Conflict as exc:
                    raise Conflict(
                        f"Duplicado AFIP en importación (fila {fila.get('fila')}): {exc.message}",
                        extra={**(exc.extra or {}), "fila": str(fila.get("fila")), "columna": "Factura"},
                    ) from exc
            self._uow.commit()
        except Exception:
            self._uow.rollback()
            raise
        return {"ok": True, "creados": creados}
