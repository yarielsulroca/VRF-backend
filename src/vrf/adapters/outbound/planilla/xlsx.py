from io import BytesIO

from vrf.domain.ports.planilla import EscritorPlanilla


class OpenpyxlEscritor(EscritorPlanilla):
    def xlsx(self, encabezados: list[str], filas: list[list[str]]) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise RuntimeError("openpyxl no está instalado") from exc

        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()

        bold = Font(bold=True)
        for col, h in enumerate(encabezados, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = bold

        for row_idx, fila in enumerate(filas, 2):
            for col_idx, val in enumerate(fila, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")

        for col_idx in range(1, len(encabezados) + 1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = max(14, len(encabezados[col_idx - 1]) + 2)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
