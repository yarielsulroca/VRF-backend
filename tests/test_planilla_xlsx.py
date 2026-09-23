from datetime import date
from io import BytesIO

from openpyxl import Workbook

from vrf.adapters.outbound.planilla.xlsx import OpenpyxlEscritor
from vrf.application.use_cases.planilla import (
    _fmt_money,
    _money,
    _parse_fecha,
    _parse_numero_factura,
    leer_filas_xlsx,
)
from vrf.domain.enums import TipoAfip


def test_xlsx_bytes_empieza_con_pk_y_tiene_una_hoja():
    escritor = OpenpyxlEscritor()
    data = escritor.xlsx(["Fecha", "Proveedor"], [["01/09/2026", "Acme"]])
    assert data[:2] == b"PK"


def test_fmt_money_argentino():
    assert _fmt_money("1234.56") == "1.234,56"
    assert _fmt_money("63550") == "63.550,00"


def test_money_acepta_float_excel():
    assert _money(15.0) is not None
    assert str(_money(15.0)) == "15.00"
    assert _money(3.15) is not None


def test_parse_fecha_espaniol():
    assert _parse_fecha("01-jul-26") == date(2026, 7, 1)


def test_parse_numero_factura_sf_y_completo():
    tipo, pto, nro, txt = _parse_numero_factura("SF")
    assert tipo == TipoAfip.NINGUNO
    assert pto is None and nro is None
    assert txt == "SF"
    tipo, pto, nro, txt = _parse_numero_factura("0001-00000027")
    assert tipo == TipoAfip.C
    assert pto == 1 and nro == 27


def test_leer_compras_xlsx_formato_libro_iva():
    wb = Workbook()
    ws = wb.active
    ws.title = "MES"
    ws["A3"] = "VRF S.A."
    ws["C3"] = "CUIT Nº 33-71117492-9"
    ws["A4"] = "LIBRO IVA COMPRAS"
    ws.append([])
    ws.append([])
    ws.append(
        [
            "FECHA",
            "NUMERO FACTURA",
            "PROVEEDOR",
            "C.U.I.T.",
            'FACTURA. "C"',
            'FACTURA. "A"',
            "IVA 0,21",
            0.105,
            "IVA 0,27",
            "IVA 10,5",
            "NO.G",
            "P.IVA",
            "P.IB",
            "TOTAL",
        ]
    )
    ws.append(
        [date(2026, 7, 1), "7333", "RELD", None, None, 15, 3.15, None, None, None, None, None, None, 18.15]
    )
    buf = BytesIO()
    wb.save(buf)
    _, filas = leer_filas_xlsx(buf.getvalue())
    assert len(filas) == 1
    assert filas[0]["proveedor"] == "RELD"
    assert filas[0]["neto_a"] == 15
    assert filas[0]["iva_21"] == 3.15
    assert filas[0]["total"] == 18.15
